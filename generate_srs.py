from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(2.5)

# ── Colour constants ──────────────────────────────────────────────────────────
DARK_BLUE  = RGBColor(0x1F, 0x36, 0x64)   # section headings
MID_BLUE   = RGBColor(0x2E, 0x75, 0xB6)   # table header bg
LIGHT_BLUE = RGBColor(0xD6, 0xE4, 0xF0)   # alternate row bg
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
BLACK      = RGBColor(0x00, 0x00, 0x00)
HEADER_TXT = RGBColor(0xFF, 0xFF, 0xFF)

# ── Helpers ───────────────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color: str):
    """Fill a table cell background."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, **kwargs):
    """Set individual borders on a cell."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = OxmlElement(f'w:{edge}')
        tag.set(qn('w:val'),   kwargs.get('val',   'single'))
        tag.set(qn('w:sz'),    kwargs.get('sz',    '4'))
        tag.set(qn('w:space'), '0')
        tag.set(qn('w:color'), kwargs.get('color', '2E75B6'))
        tcBorders.append(tag)
    tcPr.append(tcBorders)


def cell_para(cell, text, bold=False, italic=False,
              font_size=10, color=BLACK, center=False):
    """Write text into a cell paragraph with formatting."""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    para = cell.paragraphs[0]
    para.clear()
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    run.bold   = bold
    run.italic = italic
    run.font.size  = Pt(font_size)
    run.font.color.rgb = color
    return para


def heading(level: int, text: str):
    """Add a styled heading."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = DARK_BLUE
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = DARK_BLUE
    else:
        run.font.size = Pt(11)
        run.font.color.rgb = MID_BLUE
    return p


def fr_table(fr_id: str, fr_title: str, rows: list):
    """
    Render one FR block:
      FR-XX  <Title>
      Req. ID | Functional Requirements
      FR-XX-01 | ...
      ...
    """
    # Section label  e.g.  "FR-01  User Registration"
    p = doc.add_paragraph()
    run = p.add_run(f"{fr_id}  {fr_title}")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = DARK_BLUE
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(2)

    tbl = doc.add_table(rows=1 + len(rows), cols=2)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Column widths: ID col narrow, description col wide
    tbl.columns[0].width = Cm(3.5)
    tbl.columns[1].width = Cm(12.5)

    # Header row
    hdr = tbl.rows[0].cells
    set_cell_bg(hdr[0], '2E75B6')
    set_cell_bg(hdr[1], '2E75B6')
    cell_para(hdr[0], 'Req. ID',                  bold=True, color=WHITE, center=True)
    cell_para(hdr[1], 'Functional Requirements',   bold=True, color=WHITE)

    # Data rows
    for i, (req_id, req_text) in enumerate(rows):
        cells = tbl.rows[i + 1].cells
        bg = 'D6E4F0' if i % 2 == 0 else 'FFFFFF'
        set_cell_bg(cells[0], bg)
        set_cell_bg(cells[1], bg)
        cell_para(cells[0], req_id,   bold=True,  center=True)
        cell_para(cells[1], req_text, bold=False)

    doc.add_paragraph()   # breathing space


def nfr_table(nfr_id: str, nfr_title: str, rows: list):
    """Same layout but labelled Non-Functional Requirements."""
    p = doc.add_paragraph()
    run = p.add_run(f"{nfr_id}  {nfr_title}")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = DARK_BLUE
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(2)

    tbl = doc.add_table(rows=1 + len(rows), cols=2)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.columns[0].width = Cm(3.5)
    tbl.columns[1].width = Cm(12.5)

    hdr = tbl.rows[0].cells
    set_cell_bg(hdr[0], '1F3864')
    set_cell_bg(hdr[1], '1F3864')
    cell_para(hdr[0], 'Req. ID',                       bold=True, color=WHITE, center=True)
    cell_para(hdr[1], 'Non-Functional Requirements',    bold=True, color=WHITE)

    for i, (req_id, req_text) in enumerate(rows):
        cells = tbl.rows[i + 1].cells
        bg = 'EAF0F7' if i % 2 == 0 else 'FFFFFF'
        set_cell_bg(cells[0], bg)
        set_cell_bg(cells[1], bg)
        cell_para(cells[0], req_id,   bold=True,  center=True)
        cell_para(cells[1], req_text, bold=False)

    doc.add_paragraph()


def module_heading(label: str, title: str):
    """Bold letter heading  e.g.  A- Users Module"""
    p = doc.add_paragraph()
    run = p.add_run(f"{label}  {title}")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = DARK_BLUE
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(4)


# ══════════════════════════════════════════════════════════════════════════════
# DOCUMENT TITLE
# ══════════════════════════════════════════════════════════════════════════════
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_p.add_run('AskAITutor – AI-Powered Lecture Assistant')
run.bold = True
run.font.size = Pt(18)
run.font.color.rgb = DARK_BLUE

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_run = sub_p.add_run('Software Requirements Specification\nFunctional & Non-Functional Requirements')
sub_run.font.size = Pt(12)
sub_run.font.color.rgb = MID_BLUE
doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# 2.3.1  FUNCTIONAL REQUIREMENTS
# ══════════════════════════════════════════════════════════════════════════════
heading(1, '2.3.1  Functional Requirements')
doc.add_paragraph()

# ─────────────────────────────────────────────────────────────────────────────
# A – USERS MODULE
# ─────────────────────────────────────────────────────────────────────────────
module_heading('A-', 'Users Module')

fr_table('FR-01', 'User Registration', [
    ('FR-01-01', 'The system shall allow the user to enter their full name during account creation.'),
    ('FR-01-02', 'The system shall allow the user to enter their email address during account creation.'),
    ('FR-01-03', 'The system shall allow the user to enter and confirm their password during account creation.'),
    ('FR-01-04', 'The system shall validate the email format and reject malformed addresses before form submission.'),
    ('FR-01-05', 'The system shall prevent duplicate registrations and display an appropriate error when an email is already in use.'),
])

fr_table('FR-02', 'User Login', [
    ('FR-02-01', 'The system shall allow the user to log in by entering their registered email address.'),
    ('FR-02-02', 'The system shall allow the user to log in by entering their password.'),
    ('FR-02-03', 'The system shall issue a short-lived JWT access token (15 minutes) upon successful authentication.'),
    ('FR-02-04', 'The system shall store a long-lived refresh token in an HTTP-only cookie upon successful login.'),
    ('FR-02-05', 'The system shall display an error message when login credentials are incorrect.'),
])

fr_table('FR-03', 'Password Management', [
    ('FR-03-01', 'The system shall allow the user to request a password reset by submitting their registered email address.'),
    ('FR-03-02', 'The system shall send a time-bound (1-hour) password reset link to the user\'s registered email.'),
    ('FR-03-03', 'The system shall allow the user to set a new password using a valid, unexpired reset token.'),
    ('FR-03-04', 'The system shall invalidate all active refresh tokens upon a successful password reset.'),
    ('FR-03-05', 'The system shall display a uniform success message for password-reset requests regardless of whether the email exists, to prevent email enumeration.'),
])

fr_table('FR-04', 'User Interface Interaction', [
    ('FR-04-01', 'The system shall redirect unauthenticated users to the login page when they attempt to access protected routes.'),
    ('FR-04-02', 'The system shall display inline validation error messages beneath each form field.'),
    ('FR-04-03', 'The system shall show a loading indicator during API requests to provide user feedback.'),
    ('FR-04-04', 'The system shall display success and error feedback messages (toasts/alerts) following user actions.'),
    ('FR-04-05', 'The system shall automatically refresh the JWT access token in the background using the refresh token when a 401 Unauthorized response is received.'),
])

# ─────────────────────────────────────────────────────────────────────────────
# B – COURSES & LECTURES MODULE
# ─────────────────────────────────────────────────────────────────────────────
module_heading('B-', 'Courses & Lectures Module')

fr_table('FR-05', 'Browse Lectures', [
    ('FR-05-01', 'The system shall display a list of all available lectures to authenticated users on the Courses page.'),
    ('FR-05-02', 'The system shall display each lecture\'s sequential number, title, and duration.'),
    ('FR-05-03', 'The system shall display the user\'s watched progress percentage (0–100%) alongside each lecture.'),
    ('FR-05-04', 'The system shall allow the user to select any lecture from the list to begin watching.'),
    ('FR-05-05', 'The system shall visually distinguish lectures that have been fully watched from unwatched ones.'),
])

fr_table('FR-06', 'Lecture Playlist Navigation', [
    ('FR-06-01', 'The system shall display a sidebar playlist column listing all lectures in the course while viewing a lecture.'),
    ('FR-06-02', 'The system shall highlight the currently playing lecture within the playlist.'),
    ('FR-06-03', 'The system shall allow the user to switch to any other lecture directly from the playlist sidebar.'),
    ('FR-06-04', 'The system shall display the total number of lectures and how many have been completed.'),
])

# ─────────────────────────────────────────────────────────────────────────────
# C – LECTURE VIEWING MODULE
# ─────────────────────────────────────────────────────────────────────────────
module_heading('C-', 'Lecture Viewing Module')

fr_table('FR-07', 'Watch Lecture Video', [
    ('FR-07-01', 'The system shall embed and stream YouTube-hosted lecture videos on the lecture page.'),
    ('FR-07-02', 'The system shall allow the user to play, pause, seek, and adjust volume within the embedded video player.'),
    ('FR-07-03', 'The system shall display the lecture title and associated metadata (e.g., duration, lecture number) on the viewing page.'),
    ('FR-07-04', 'The system shall provide a navigation button allowing the user to proceed to the next lecture directly from the viewing page.'),
])

# ─────────────────────────────────────────────────────────────────────────────
# D – AI CHAT MODULE
# ─────────────────────────────────────────────────────────────────────────────
module_heading('D-', 'AI Chat Module')

fr_table('FR-08', 'Text-Based Question & Answer', [
    ('FR-08-01', 'The system shall provide a slide-up chat drawer on the lecture page allowing the user to type questions about the lecture.'),
    ('FR-08-02', 'The system shall send the user\'s typed question to the AI tutor for processing upon submission.'),
    ('FR-08-03', 'The system shall display the AI tutor\'s response in the chat interface in a visually distinct message bubble.'),
    ('FR-08-04', 'The system shall scope AI tutor responses strictly to the content of the current lecture to prevent off-topic or hallucinated answers.'),
    ('FR-08-05', 'The system shall display suggestion chips (quick-question shortcuts) to help the user formulate lecture-related questions.'),
    ('FR-08-06', 'The system shall display a typing indicator animation while the AI tutor is generating a response.'),
])

fr_table('FR-09', 'Voice-Based Question & Answer', [
    ('FR-09-01', 'The system shall allow the user to activate voice input by pressing a microphone button in the chat interface.'),
    ('FR-09-02', 'The system shall convert the user\'s spoken question to text in real time using the browser\'s Web Speech Recognition API.'),
    ('FR-09-03', 'The system shall automatically submit the transcribed voice question after detecting 1.5 seconds of silence.'),
    ('FR-09-04', 'The system shall read the AI tutor\'s response aloud using Text-to-Speech synthesis after the answer is received.'),
    ('FR-09-05', 'The system shall allow the user to pause or stop the AI\'s spoken response at any time.'),
    ('FR-09-06', 'The system shall detect whether the user\'s browser supports voice features and display an informative message if voice is unavailable.'),
])

fr_table('FR-10', 'Chat Interface Management', [
    ('FR-10-01', 'The system shall allow the user to toggle the chat drawer open and closed via a visible toggle bar.'),
    ('FR-10-02', 'The system shall display chat messages in chronological order with timestamps.'),
    ('FR-10-03', 'The system shall visually differentiate user messages from AI tutor responses using distinct colours and alignment.'),
    ('FR-10-04', 'The system shall allow the user to switch between text input mode and voice input mode within the chat interface.'),
    ('FR-10-05', 'The system shall display the full chat history for the current lecture session within the chat drawer.'),
])

# ══════════════════════════════════════════════════════════════════════════════
# 2.3.2  NON-FUNCTIONAL REQUIREMENTS
# ══════════════════════════════════════════════════════════════════════════════
doc.add_page_break()
heading(1, '2.3.2  Non-Functional Requirements')
doc.add_paragraph()

nfr_table('NFR-01', 'Performance', [
    ('NFR-01-01', 'The system shall return the AI tutor\'s response within 3–5 seconds of the user submitting a question.'),
    ('NFR-01-02', 'The system shall load the lecture page, including the embedded video player and chat drawer, within 3 seconds on a standard broadband connection.'),
    ('NFR-01-03', 'The system shall support multiple concurrent authenticated user sessions without degradation in response time or functionality.'),
    ('NFR-01-04', 'The system shall activate the speech recognition listener within 1 second of the user pressing the microphone button.'),
    ('NFR-01-05', 'The system shall complete user authentication (login/register) API calls within 2 seconds under normal load.'),
])

nfr_table('NFR-02', 'Security', [
    ('NFR-02-01', 'The system shall authenticate every protected API request using a valid JWT access token provided in the Authorization header.'),
    ('NFR-02-02', 'The system shall store refresh tokens exclusively in HTTP-only, Secure, SameSite=Strict cookies to mitigate XSS and CSRF attacks.'),
    ('NFR-02-03', 'The system shall hash all refresh tokens and password-reset tokens using SHA-256 before persisting them in the database.'),
    ('NFR-02-04', 'The system shall hash user passwords using bcrypt with a minimum of 12 salt rounds before storage.'),
    ('NFR-02-05', 'The system shall enforce rate limiting on all authentication endpoints, allowing no more than 10 requests per 15 minutes per IP address.'),
    ('NFR-02-06', 'The system shall invalidate all existing refresh tokens after a successful password reset to prevent unauthorized session continuation.'),
    ('NFR-02-07', 'The system shall apply security headers (via Helmet.js) including X-Frame-Options: DENY and X-Content-Type-Options: nosniff to all responses.'),
    ('NFR-02-08', 'The system shall restrict CORS to the whitelisted frontend origin only.'),
])

nfr_table('NFR-03', 'Accuracy', [
    ('NFR-03-01', 'The AI tutor shall respond only with information derived from the content of the currently active lecture, avoiding hallucinated or out-of-scope answers.'),
    ('NFR-03-02', 'The voice-to-text transcription shall achieve an accuracy of at least 85% for clearly spoken English under normal acoustic conditions.'),
    ('NFR-03-03', 'The system shall handle unrecognised speech, empty inputs, and ambiguous queries gracefully by prompting the user rather than crashing.'),
    ('NFR-03-04', 'The system shall consistently produce the same AI response for identical questions about the same lecture content.'),
])

nfr_table('NFR-04', 'Usability', [
    ('NFR-04-01', 'The system shall provide a clean, intuitive interface that requires no technical training for students to use effectively.'),
    ('NFR-04-02', 'The system shall be fully operable on laptops, tablets, and desktop computers with varying screen resolutions.'),
    ('NFR-04-03', 'The system shall provide clear visual feedback (loading spinners, inline errors, success messages) for every user-initiated action.'),
    ('NFR-04-04', 'The system shall offer suggestion chips to guide users who are unsure how to phrase a question about the lecture.'),
    ('NFR-04-05', 'The chat drawer shall remain accessible at all times during lecture playback without obscuring the video player.'),
])

nfr_table('NFR-05', 'Maintainability', [
    ('NFR-05-01', 'The system shall follow a modular, layered architecture that separates the frontend UI, backend API, and database layers, allowing each to be independently updated.'),
    ('NFR-05-02', 'The backend shall apply the service-layer pattern, keeping business logic in service files distinct from route controllers.'),
    ('NFR-05-03', 'All system modules, API endpoints, and custom hooks shall be documented to facilitate future maintenance and developer onboarding.'),
    ('NFR-05-04', 'The frontend shall use reusable UI components (Button, Input, ChatDrawer, etc.) to minimise code duplication across pages.'),
])

nfr_table('NFR-06', 'Reliability', [
    ('NFR-06-01', 'The system shall maintain a minimum uptime of 99% to ensure uninterrupted access for students during their study sessions.'),
    ('NFR-06-02', 'The system shall automatically log all errors and unhandled exceptions using a structured logging library (Winston) without disrupting active user sessions.'),
    ('NFR-06-03', 'The system shall handle network failures, API timeouts, and unexpected server errors gracefully by displaying informative messages to the user.'),
    ('NFR-06-04', 'The system shall recover automatically from transient database connection failures through connection-pool retry mechanisms.'),
])

nfr_table('NFR-07', 'Scalability', [
    ('NFR-07-01', 'The system shall use PostgreSQL connection pooling to efficiently manage increasing numbers of concurrent database connections.'),
    ('NFR-07-02', 'The backend architecture shall support horizontal scaling (multiple server instances) to accommodate a growing user base without architectural changes.'),
    ('NFR-07-03', 'The frontend build shall be optimised via Next.js static generation and code-splitting to reduce initial load times as the content library grows.'),
])

# ══════════════════════════════════════════════════════════════════════════════
# SAVE
# ══════════════════════════════════════════════════════════════════════════════
out = '/home/kiran-shehzadi/askaitutor/AskAITutor_FRs_NFRs.docx'
doc.save(out)
print(f'Saved → {out}')
